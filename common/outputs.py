"""Machine-readable and human-readable output writers.

Supports the CLI --format contract: json | csv | md | xlsx | pdf.
"""

from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any, Iterable, Optional

from .errors import UnsupportedFileError

FORMATS = ("json", "csv", "md", "xlsx", "pdf")


def write_json(path: str | Path, obj: Any) -> str:
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    with open(p, "w", encoding="utf-8") as fh:
        json.dump(obj, fh, indent=2, ensure_ascii=False)
    return str(p)


def write_csv(path: str | Path, rows: Iterable[dict], columns: Optional[list[str]] = None) -> str:
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    rows = list(rows)
    if not rows:
        rows = [{}]
    cols = columns or list(rows[0].keys())
    with open(p, "w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=cols, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({c: row.get(c) for c in cols})
    return str(p)


def write_markdown(path: str | Path, sections: Iterable[tuple[str, str]]) -> str:
    """Write markdown from an ordered list of (heading, body) sections."""
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    lines: list[str] = []
    for heading, body in sections:
        if heading:
            lines.append(f"## {heading}\n")
        lines.append(f"{body}\n")
    with open(p, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines))
    return str(p)


def _table_to_markdown(rows: list[dict]) -> str:
    if not rows:
        return "_no rows_"
    cols = list(rows[0].keys())
    header = "| " + " | ".join(cols) + " |"
    sep = "| " + " | ".join(["---"] * len(cols)) + " |"
    body = []
    for row in rows:
        body.append("| " + " | ".join(str(row.get(c, "")) for c in cols) + " |")
    return "\n".join([header, sep, *body])


def write_xlsx(path: str | Path, sheets: dict[str, list[dict]] | None = None) -> str:
    """Write an xlsx workbook; sheets = {sheet_name: list-of-dict rows}."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError as exc:  # pragma: no cover
        raise UnsupportedFileError(
            "xlsx output requires openpyxl (pip install openpyxl)."
        ) from exc

    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    sheets = sheets or {"Sheet1": []}
    for name, rows in sheets.items():
        ws = wb.create_sheet(title=name[:31])
        rows = list(rows) or [{}]
        cols = list(rows[0].keys())
        for c, col in enumerate(cols, start=1):
            cell = ws.cell(row=1, column=c, value=col)
            cell.font = Font(bold=True)
        for r, row in enumerate(rows, start=2):
            for c, col in enumerate(cols, start=1):
                ws.cell(row=r, column=c, value=row.get(col))
    wb.save(p)
    return str(p)


def write_pdf(path: str | Path, title: str, sections: Iterable[tuple[str, str]]) -> str:
    """Write a simple PDF report (requires reportlab)."""
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import inch
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    except ImportError as exc:  # pragma: no cover
        raise UnsupportedFileError(
            "pdf output requires reportlab (pip install reportlab)."
        ) from exc

    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(str(p), pagesize=letter)
    title_style = ParagraphStyle("title", fontSize=16, leading=20, spaceAfter=12)
    head_style = ParagraphStyle("head", fontSize=12, leading=16, spaceBefore=10, spaceAfter=4)
    body_style = ParagraphStyle("body", fontSize=9, leading=12)
    story = [Paragraph(title, title_style)]
    for heading, body in sections:
        if heading:
            story.append(Paragraph(heading, head_style))
        for line in str(body).splitlines():
            story.append(Paragraph(line.replace("|", " / "), body_style))
        story.append(Spacer(1, 6))
    doc.build(story)
    return str(p)


def write_output(path: str | Path, fmt: str, data: Any) -> str:
    """Dispatch a write by format.

    data conventions:
      json -> any serializable object
      csv  -> list[dict]
      md   -> list[(heading, body)]  (body may be a table string)
      xlsx -> {"sheet": [dict,...], ...}
      pdf  -> (title, list[(heading, body)])
    """
    fmt = (fmt or "json").lower()
    if fmt not in FORMATS:
        raise UnsupportedFileError(f"Unsupported output format: {fmt}")
    if fmt == "json":
        return write_json(path, data)
    if fmt == "csv":
        return write_csv(path, data)
    if fmt == "md":
        return write_markdown(path, data)
    if fmt == "xlsx":
        return write_xlsx(path, data)
    title, sections = data
    return write_pdf(path, title, sections)
